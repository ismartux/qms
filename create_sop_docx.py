import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for a table cell (values in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Set subtle borders for a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, title, text, box_type="info"):
    """Add a beautifully styled callout box with a left colored border."""
    # Palette
    configs = {
        "info": {"bg": "F0F9FF", "border": "0284C7", "icon": "ℹ️", "title_color": RGBColor(2, 132, 199)},
        "tip": {"bg": "F0FDF4", "border": "16A34A", "icon": "💡", "title_color": RGBColor(22, 163, 74)},
        "warning": {"bg": "FFFBEB", "border": "D97706", "icon": "⚠️", "title_color": RGBColor(217, 119, 6)},
        "important": {"bg": "FEF2F2", "border": "DC2626", "icon": "❗", "title_color": RGBColor(220, 38, 38)},
        "step": {"bg": "F8FAFC", "border": "3B82F6", "icon": "🎯", "title_color": RGBColor(30, 64, 175)}
    }
    cfg = configs.get(box_type, configs["info"])
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, cfg["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{cfg["border"]}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    run_title = p.add_run(f"{cfg['icon']} {title}\n")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = cfg["title_color"]
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(51, 65, 85) # slate-700
    
    # Spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(2)

def build_sop():
    doc = docx.Document()
    
    # Page setup - Standard Letter, 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_foot.text = "QIsmartuX QMS — Standard Operating Procedure | Page "
        p_foot.runs[0].font.name = "Calibri"
        p_foot.runs[0].font.size = Pt(8.5)
        p_foot.runs[0].font.color.rgb = RGBColor(148, 163, 184)

    # Color Palette Constants
    NAVY = RGBColor(15, 23, 42)       # Primary headings (#0F172A)
    ACCENT_BLUE = RGBColor(2, 132, 199) # Section numbers / highlights (#0284C7)
    SLATE_BODY = RGBColor(51, 65, 85)  # Body text (#334155)
    MUTED_GRAY = RGBColor(100, 116, 139) # (#64748B)

    # Helper styling functions
    def add_sop_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(15, 23, 42)

    def add_sop_subtitle(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(18)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(2, 132, 199)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(15)
        run.font.color.rgb = NAVY

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(12.5)
        run.font.color.rgb = ACCENT_BLUE

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(30, 41, 59)

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(10.5)
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = SLATE_BODY
        return p

    def add_bullet(text, bold_prefix=None, level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(10.5)
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = SLATE_BODY
        return p

    def add_step(step_number, step_title, description):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        # Step Badge
        badge = p.add_run(f"STEP {step_number}: ")
        badge.bold = True
        badge.font.name = "Calibri"
        badge.font.size = Pt(10.5)
        badge.font.color.rgb = ACCENT_BLUE
        
        title_run = p.add_run(f"{step_title}\n")
        title_run.bold = True
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(10.5)
        title_run.font.color.rgb = NAVY
        
        desc_run = p.add_run(description)
        desc_run.font.name = "Calibri"
        desc_run.font.size = Pt(10)
        desc_run.font.color.rgb = SLATE_BODY

    # =========================================================================
    # DOCUMENT COVER / HEADER
    # =========================================================================
    add_sop_title("STANDARD OPERATING PROCEDURE (SOP)")
    add_sop_subtitle("QIsmartuX Quality Management System (QMS) — End-User Operations Manual")

    # Document Control Metadata Table
    meta_table = doc.add_table(rows=5, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, "CBD5E1", "4")
    
    col_widths = [Inches(1.5), Inches(1.8), Inches(1.4), Inches(1.8)]
    for row in meta_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    headers = [
        ("Document ID:", "SOP-QMS-OPS-2026-01", "Effective Date:", "August 2026"),
        ("System Name:", "QIsmartuX QMS Platform", "Version:", "2.0 (Live Production)"),
        ("Target Users:", "Operators, Supervisors, PQE, Safety, Managers", "Department:", "Quality, Operations & EHS"),
        ("Access Level:", "Role-Based (Plant & Line Scoped)", "Review Cycle:", "Annual / Revision Triggered"),
        ("Document Owner:", "Quality Assurance & Operational Excellence", "Approved By:", "Head of Plant Quality")
    ]

    for row_idx, data in enumerate(headers):
        for col_idx in range(4):
            cell = meta_table.cell(row_idx, col_idx)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(data[col_idx])
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            if col_idx % 2 == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(71, 85, 105)
                set_cell_background(cell, "F1F5F9")
            else:
                run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_callout_box(
        doc,
        "Purpose & Non-Technical User Objective",
        "This Standard Operating Procedure (SOP) is designed specifically for plant end-users. It explains, in simple, practical, step-by-step instructions, how to log in, choose your active work shift, conduct quality inspections on shop-floor lines, upload defect photos, submit approvals, manage corrective actions (CAPA), perform safety audits, and monitor live metrics.",
        "info"
    )

    # =========================================================================
    # TABLE OF CONTENTS SUMMARY
    # =========================================================================
    add_h1("Table of Contents")
    toc_items = [
        "1. System Overview & User Role Responsibilities",
        "2. System Access & How to Log In",
        "3. Setting Up & Activating Your Work Session (Shop, Line, Shift, Model)",
        "4. Performing Quality Inspections (Operator Step-by-Step Guide)",
        "5. Dynamic Forms & Parameterized Inspections",
        "6. Reviewing Submissions & Inspection History",
        "7. Review & Approval Workflow (Supervisor & PQE Guide)",
        "8. Corrective & Preventive Actions (CAPA Management)",
        "9. Environment, Health & Safety (EHS) Audits",
        "10. Real-Time Dashboards & Analytics Guide",
        "11. Inspection Scheduler & Shift Routine Triggers",
        "12. Form Builder Guide (Checklist & Dynamic Form Setup)",
        "13. Plant Administration & User Management",
        "14. Offline Mode & Auto-Sync Guidelines",
        "15. Best Practices, Do's & Don'ts, FAQs and Troubleshooting"
    ]
    for item in toc_items:
        add_bullet(item)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECTION 1: SYSTEM OVERVIEW & USER ROLES
    # =========================================================================
    add_h1("1. System Overview & User Role Responsibilities")
    add_body(
        "QIsmartuX QMS is a digital manufacturing quality and safety execution platform designed for production shop floors. "
        "It replaces manual paper checklists, provides real-time defect notifications, enforces mandatory photographic evidence for non-conformances, "
        "and coordinates multi-stage managerial approvals across all production lines."
    )

    add_h2("1.1 User Roles & What You Can Do")
    
    role_table = doc.add_table(rows=6, cols=3)
    role_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    role_table.autofit = False
    set_table_borders(role_table, "CBD5E1", "4")
    
    r_widths = [Inches(1.6), Inches(2.2), Inches(2.7)]
    for row in role_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = r_widths[i]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # Header Row
    headers_role = ["Role Name", "Primary Responsibilities", "System Features Accessed"]
    for i, h in enumerate(headers_role):
        cell = role_table.cell(0, i)
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    roles_data = [
        ("Line Operator\n(IPQC / OQC / FQC)", "Conducts routine quality checks on production lines, measures parameters, logs defects, uploads proof photos, and submits records.", "Work Session Setup, Quality Checklists, Dynamic Forms, Operator Dashboard."),
        ("Line Leader /\nShop Supervisor", "Oversees line execution, monitors hourly/shift inspection completion, tracks defect rates, and conducts initial stage verification.", "Supervisor Dashboard, Work Session Monitoring, Pending Approvals, Scheduled Checklists."),
        ("PQE / Quality Engineer", "Reviews inspection records, evaluates defect severity, approves or rejects submissions with remarks, and raises CAPA tickets.", "Pending Approvals, Quality Review Popup, CAPA Management, Inspection Reports."),
        ("CAPA Lead /\nAction Owner", "Investigates root causes (5-Why Analysis), implements corrective and preventive actions, uploads evidence, and marks actions done.", "CAPA Dashboard, Assigned CAPA Workspace, Action Evidence Upload."),
        ("EHS Auditor /\nSafety Officer", "Performs daily safety rounds, conducts hazard assessments, scans station QR codes, and calculates safety risk scores.", "EHS Dashboard, Daily Safety Checklists, Special Incident Audits, Hazard Reports.")
    ]

    for row_idx, r_data in enumerate(roles_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(3):
            cell = role_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(r_data[col_idx])
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = SLATE_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECTION 2: SYSTEM ACCESS & HOW TO LOG IN
    # =========================================================================
    add_h1("2. System Access & How to Log In")
    add_body("The QMS platform can be accessed from any desktop computer, plant tablet, workstation terminal, or mobile device connected to the factory network.")

    add_step("1", "Open Your Web Browser", "Open Google Chrome, Microsoft Edge, or Safari on your workstation, tablet, or smartphone.")
    add_step("2", "Navigate to the Application URL", "Type your plant QMS web address in the address bar (e.g., http://qms.internal.plant/ or your designated server address) and press Enter.")
    add_step("3", "Enter Your Credentials", "On the login screen, enter your assigned Username (or Employee ID) and your Password.")
    add_step("4", "Click the [Sign In] Button", "Click the blue 'Sign In' button to authenticate. Upon successful login, the system will automatically direct you to your role-specific Home Screen.")

    add_callout_box(
        doc,
        "First-Time Login & Security Tip",
        "If this is your first time logging in, your supervisor will provide a temporary password. Remember to keep your credentials confidential and never share accounts. Always log out when stepping away from a shared shop-floor tablet.",
        "tip"
    )

    add_h2("2.1 Understanding the Main Home Navigation")
    add_body("Once logged in, the top navigation header and home cards display only the sections you are authorized to access:")
    add_bullet("Quality Management (IPQC): For accessing inspection forms, live quality checklists, and operator metrics.", "• ")
    add_bullet("Environment, Health & Safety (EHS): For performing daily safety walks and incident checklists.", "• ")
    add_bullet("CAPA Management: For viewing and responding to assigned corrective and preventive actions.", "• ")
    add_bullet("Pending Approvals: For supervisors and PQE engineers to verify completed inspection records.", "• ")
    add_bullet("Dashboards & Insights: For viewing live operational charts, trends, and compliance summaries.", "• ")

    # =========================================================================
    # SECTION 3: WORK SESSION / CONTEXT SETUP
    # =========================================================================
    add_h1("3. Setting Up & Activating Your Work Session")
    add_body(
        "Before an operator or inspector can fill out any quality checklist, an active 'Work Session' (Work Context) must be selected. "
        "This ensures all inspection data, measurements, and photos are automatically tagged to the correct Shop, Production Line, Model, and Shift."
    )

    add_h2("3.1 Step-by-Step: Creating & Activating a Work Session")
    add_step("1", "Click [Change Session] or [Create Work Session]", "On the home screen or quality checklist page, locate the session banner at the top and click 'Change Session' or 'Create Work Session'.")
    add_step("2", "Select Your Shop / Area", "Click the 'Shop / Area' dropdown and select the active manufacturing shop (e.g., SMT Shop, Assembly Shop, Packaging Area).")
    add_step("3", "Select the Production Line", "Click the 'Line' dropdown to choose your specific production line (e.g., Line 01, Line 02, SMT-A).")
    add_step("4", "Select the Product / Model", "Choose the active model currently being manufactured on your line from the 'Product / Model' searchable list.")
    add_step("5", "Select the Model Color (If Applicable)", "Choose the product color variant running on the line (e.g., Midnight Black, Glacier Blue).")
    add_step("6", "Select the Shift & Date", "Select your active shift (e.g., Shift A, Shift B, General Shift, Night Shift) and verify the production date.")
    add_step("7", "Click [Activate Session]", "Click the primary button to save and activate your context. The top bar will now show your active Line, Product, and Shift badge.")

    add_callout_box(
        doc,
        "Shift Handover & Model Changeover Rule",
        "Whenever a new shift begins or the production line changes to a different model, the operator MUST activate a new Work Session before recording inspections. This prevents mixing inspection data between models.",
        "warning"
    )

    # =========================================================================
    # SECTION 4: PERFORMING QUALITY INSPECTIONS (OPERATOR GUIDE)
    # =========================================================================
    add_h1("4. Performing Quality Inspections (Operator Step-by-Step Guide)")
    add_body(
        "The Quality Inspection engine presents operators with a clean, step-by-step digital checklist. "
        "It prevents common mistakes by automatically validating required answers and enforcing photographic evidence when defects occur."
    )

    add_h2("4.1 Starting an Inspection Checklist")
    add_step("1", "Navigate to Quality Checklists", "Click 'IPQC Inspections' or 'Forms' from your navigation bar or home screen.")
    add_step("2", "Select the Assigned Checklist Template", "You will see cards showing the inspection forms available for your active line (e.g., 'IPQC Daily Station Inspection', 'SMT First Piece Inspection'). Click on the form card to open it.")
    add_step("3", "Review Checklist Header Info", "The top header displays the Section Title, ISO Form Number, Template Version, and your current Progress (0% to 100%).")

    add_h2("4.2 Answering Inspection Items Step-by-Step")
    add_body("Each question item on the form will have one of the following input types:")

    add_bullet("Pass / Fail / NA (Radio Cards): Tap [Pass] (Green) if the item meets quality standards, [Fail] (Red) if there is a defect or non-conformance, or [NA] (Gray) if the item does not apply to this model.", "1. Pass / Fail Status: ")
    add_bullet("Numerical Measurements: Type the exact measured value (e.g., torque reading 1.45 Nm, solder temperature 260°C). The system highlights out-of-spec values automatically.", "2. Number / Value Entry: ")
    add_bullet("Text Answers: Type inspection observations, batch numbers, or serial numbers in the text box.", "3. Text Inputs: ")
    add_bullet("Dropdown Selections: Tap to select standard inspection options from a predefined list.", "4. Dropdown Pickers: ")
    add_bullet("Photo Evidence Upload: Tap the camera icon to capture a live photo using your tablet/phone camera or choose an image file from your device.", "5. Photo Upload: ")

    add_h2("4.3 Handling Defects & Mandatory Rules (When an item Fails)")
    add_body("When you select [Fail] on any inspection point, the system triggers built-in quality rules:")
    add_bullet("Mandatory Defect Remark Box: A text box immediately appears below the item asking 'Explain why this failed...'. You MUST enter a clear description of the defect.", "• ")
    add_bullet("Mandatory Proof Photo: A photo capture card appears marked 'Proof Photo Required'. Tap the camera icon and take a clear picture of the physical defect.", "• ")

    add_callout_box(
        doc,
        "Zero-Defect Rule Enforcement",
        "The system will NOT permit submitting an inspection with a 'Fail' status unless both a descriptive remark and a clear defect photograph are provided. This ensures immediate clarity for supervisors and quality engineers.",
        "important"
    )

    add_h2("4.4 Navigating Sections & Final Submission")
    add_step("1", "Review Section Items", "Complete all required items on the current step. Required fields with missing answers are marked with a red warning badge.")
    add_step("2", "Click [Next Step]", "Tap the blue 'Next Step' button at the bottom of the screen to advance to the next inspection section.")
    add_step("3", "Click [Back] if Review is Needed", "You can click 'Back' at any time to re-check or edit previous section responses.")
    add_step("4", "Final Step: Click [Submit Inspection]", "On the last section, the button changes to a green 'Submit Inspection' button. Tap it to send your inspection for review.")
    add_step("5", "Confirmation Screen", "A green success banner will appear confirming 'Inspection Submitted Successfully'. The record is now locked and submitted to the supervisor/PQE queue.")

    # =========================================================================
    # SECTION 5: DYNAMIC FORMS & PARAMETERIZED INSPECTIONS
    # =========================================================================
    add_h1("5. Dynamic Forms & Parameterized Inspections")
    add_body(
        "For complex processes requiring multi-sample dimensions, custom formulas, or searchable component lookups, "
        "the QMS provides the 'Dynamic Forms Engine'."
    )

    add_h2("5.1 How Dynamic Forms Differ from Standard Checklists")
    add_bullet("Multi-Sample Recording: Allows recording measurements across Sample 1, Sample 2, Sample 3... up to Sample 10 in a single screen.", "• ")
    add_bullet("Dynamic Searchable Dropdowns: Quickly type part numbers, supplier names, or defect codes to filter long lists instantly.", "• ")
    add_bullet("Instant Standard Limit Check: Real-time visual indicators show whether each sample measurement falls within Upper Spec Limit (USL) and Lower Spec Limit (LSL).", "• ")

    add_h2("5.2 Step-by-Step: Filling a Dynamic Form")
    add_step("1", "Open Dynamic Form", "Select the dynamic form from your assigned forms list.")
    add_step("2", "Enter Header Details", "Select sample lot size, machine number, or tool cavity if prompted.")
    add_step("3", "Enter Sample Measurements", "Type in each sample's dimension. If a reading is outside tolerance, the input box highlights in amber/red alert.")
    add_step("4", "Attach Inspection Evidence", "Upload overall inspection photos or component batch labels.")
    add_step("5", "Click [Submit Dynamic Form]", "Review all entered values and tap the green submit button.")

    # =========================================================================
    # SECTION 6: REVIEWING SUBMISSIONS & INSPECTION HISTORY
    # =========================================================================
    add_h1("6. Reviewing Submissions & Inspection History")
    add_body("Operators, supervisors, and quality personnel can review all past inspection records at any time.")

    add_h2("6.1 How to Search & Filter Past Submissions")
    add_step("1", "Go to Submissions", "Click 'Submissions' in the main navigation menu.")
    add_step("2", "Use Filter Options", "Filter records by: Date Range (Today, This Week, Custom Date), Production Line, Product Model, and Approval Status (Pending, Approved, Rejected).")
    add_step("3", "View Summary List", "Each record in the list displays: Form Name, Line Code, Product, Submission Timestamp, Submitter Name, and Status Badge.")
    add_step("4", "Open Submission Details Popup", "Click the 'View' or 'Review Details' button on any record to open the full inspection report.")

    add_h2("6.2 What is Inside the Submission Detail Popup?")
    add_bullet("Complete ISO & Plant Header (Plant, Shop, Line, Product, Shift, Inspector Name, Date & Time).", "• ")
    add_bullet("Section-by-Section breakdown of all questions with operator answers.", "• ")
    add_bullet("High-resolution photo evidence previews (tap any image thumbnail to expand full-size).", "• ")
    add_bullet("Approval Timeline Tracker showing who reviewed the form, approval timestamp, and approver remarks.", "• ")
    add_bullet("Print / PDF Export button for audit documentation and customer reporting.", "• ")

    # =========================================================================
    # SECTION 7: REVIEW & APPROVAL WORKFLOW (SUPERVISOR & PQE GUIDE)
    # =========================================================================
    add_h1("7. Review & Approval Workflow (Supervisor & PQE Guide)")
    add_body(
        "To maintain strict quality assurance, inspection submissions move through a multi-stage approval workflow. "
        "Supervisors and Product Quality Engineers (PQE) must review submissions and either Approve or Reject them."
    )

    add_h2("7.1 Accessing Pending Approvals")
    add_step("1", "Navigate to Pending Approvals", "Click 'Pending Approvals' on your home screen or top navigation bar.")
    add_step("2", "View Pending Cards", "Each card shows the inspection template, submitted time, operator name, and the multi-role approval progression badge.")
    add_step("3", "Click [Review Details]", "Tap the blue 'Review Details' button on the submission card to inspect the record.")

    add_h2("7.2 Step-by-Step: Approving an Inspection Submission")
    add_step("1", "Examine All Inspection Responses", "Scroll through the inspection points. Pay special attention to any items highlighted in red or amber.")
    add_step("2", "Verify Uploaded Defect Photos & Measurements", "Click on attached photos to ensure defects are correctly identified and documented.")
    add_step("3", "Verify Approver Name", "In the approval section at the bottom, your name will be auto-filled. Ensure it matches your credentials.")
    add_step("4", "Click [Approve]", "Click the green 'Approve' button. The system updates the submission status to 'Approved' and moves it to the next workflow stage (or marks it fully approved).")

    add_h2("7.3 Step-by-Step: Rejecting an Inspection Submission")
    add_step("1", "Identify the Problem", "If an inspection has missing data, incorrect photos, or an unresolved critical defect, click [Reject].")
    add_step("2", "Enter Mandatory Rejection Reason", "A text area will open labeled 'Enter rejection reason...'. Type a detailed explanation of why the record was rejected (e.g., 'Torque readings incomplete for Station 3; re-measurement required').")
    add_step("3", "Click [Confirm Reject]", "Click the dark red 'Confirm Reject' button. The operator and line supervisor will be immediately notified to take corrective action.")

    add_callout_box(
        doc,
        "Mobile & Public Approval Links",
        "When managers or lead engineers receive inspection alert notifications on mobile devices or Lark/Feishu, they can click the secure review link to approve or reject submissions directly from their mobile browser without logging into a workstation.",
        "tip"
    )

    # =========================================================================
    # SECTION 8: CORRECTIVE & PREVENTIVE ACTIONS (CAPA)
    # =========================================================================
    add_h1("8. Corrective & Preventive Actions (CAPA Management)")
    add_body(
        "When recurring defects, critical non-conformances, or audit failures occur, the CAPA module provides a structured, "
        "closed-loop problem-solving mechanism to prevent recurrence."
    )

    add_h2("8.1 The 4-Stage CAPA Lifecycle")
    
    capa_table = doc.add_table(rows=5, cols=3)
    capa_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    capa_table.autofit = False
    set_table_borders(capa_table, "CBD5E1", "4")
    
    c_widths = [Inches(1.5), Inches(2.2), Inches(2.8)]
    for row in capa_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = c_widths[i]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # Header Row
    headers_capa = ["Stage", "Status Badge", "Action Required"]
    for i, h in enumerate(headers_capa):
        cell = capa_table.cell(0, i)
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    capa_lifecycle = [
        ("1. Issue Raised", "OPEN (Yellow)", "Non-conformance identified from failed inspection or manual safety report."),
        ("2. Assigned", "ASSIGNED (Blue)", "PQE assigns RCA Owner, CAPA Action Owner, and sets Target Due Date."),
        ("3. Action Done", "ACTION DONE (Purple)", "Assignee conducts 5-Why RCA, defines corrective/preventive plans, attaches proof, and submits."),
        ("4. Closure / Review", "CLOSED (Green)", "Quality Manager / PQE verifies action effectiveness and officially closes the CAPA.")
    ]

    for row_idx, c_data in enumerate(capa_lifecycle, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(3):
            cell = capa_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(c_data[col_idx])
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = SLATE_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("8.2 Step-by-Step: Assigning a CAPA (For Quality Lead / PQE)")
    add_step("1", "Open CAPA List", "Click 'CAPA Management' in the main navigation menu.")
    add_step("2", "Select Open CAPA", "Click on any CAPA ticket in 'OPEN' status.")
    add_step("3", "Assign RCA Owner & Action Role", "Select the department lead or manufacturing engineer responsible for Root Cause Analysis.")
    add_step("4", "Set Due Date", "Choose the target completion deadline based on issue severity (e.g., 24 hours for Critical, 3 days for Major).")
    add_step("5", "Click [Assign CAPA]", "The assignee receives an automated notification with ticket details.")

    add_h2("8.3 Step-by-Step: Executing RCA & Submitting Action Plan (For Assignee)")
    add_step("1", "Open 'My Assigned CAPAs'", "Navigate to CAPA > My Assigned CAPAs and open your assigned ticket.")
    add_step("2", "Perform Root Cause Analysis (RCA)", "Type the root cause findings in the RCA description box (e.g., 'Tool calibration drifted after 500 cycles; lack of automatic torque cutoff').")
    add_step("3", "Document Immediate Corrective Action", "Explain the immediate containment action taken (e.g., 'Quarantined 40 suspect units, recalibrated driver bit immediately').")
    add_step("4", "Document Permanent Preventive Action", "Detail the systemic fix to prevent recurrence (e.g., 'Updated PM schedule to recalibrate tool every 200 cycles; installed digital torque monitor').")
    add_step("5", "Upload Action Proof Photo / Document", "Attach verification photos or test reports proving the fix is in place.")
    add_step("6", "Click [Mark Action Done]", "Submit your completed action plan for quality review.")

    add_h2("8.4 Step-by-Step: Verification & Final Closure (For Quality Lead)")
    add_step("1", "Review Implemented Action", "Open the CAPA ticket in 'ACTION DONE' status and inspect the evidence.")
    add_step("2", "Verify Effectiveness on the Shop Floor", "Confirm that the line is running defect-free post-implementation.")
    add_step("3", "Click [Approve & Close CAPA]", "The ticket is marked 'CLOSED' and archived in the plant compliance history.")

    # =========================================================================
    # SECTION 9: ENVIRONMENT, HEALTH & SAFETY (EHS) AUDITS
    # =========================================================================
    add_h1("9. Environment, Health & Safety (EHS) Audits")
    add_body(
        "The EHS module empowers safety officers and floor supervisors to conduct structured safety audits, "
        "track emergency preparedness, verify PPE compliance, and calculate safety risk scores."
    )

    add_h2("9.1 Step-by-Step: Conducting Daily Safety Audits")
    add_step("1", "Open EHS Dashboard", "Click 'EHS Home' or 'EHS' from the main navigation.")
    add_step("2", "Select Audit Type", "Choose 'Daily Forms' for routine shop floor safety walks or 'Special Forms' for incident investigations.")
    add_step("3", "Scan Station QR Code (If Applicable)", "If auditing a designated safety zone or fire extinguisher station, tap 'Scan QR Code' to automatically log the physical station location.")
    add_step("4", "Answer Safety Audit Checklist", "Inspect and record compliance on: PPE usage, fire extinguisher accessibility, chemical storage, electrical safety, machine guarding, and slip/trip hazards.")
    add_step("5", "Log Safety Hazards & Attach Photos", "If a hazard is observed (e.g., blocked fire exit, oil spill), select 'Hazard Detected', upload a photo, and assign a hazard severity rating.")
    add_step("6", "Submit Safety Report", "Tap 'Submit Safety Audit'. The system instantly calculates the Station Risk Score and generates safety alerts if high-risk items are flagged.")

    # =========================================================================
    # SECTION 10: REAL-TIME DASHBOARDS & ANALYTICS GUIDE
    # =========================================================================
    add_h1("10. Real-Time Dashboards & Analytics Guide")
    add_body("The QMS provides role-tailored dashboards to give every team member the exact operational visibility they need:")

    add_h2("10.1 Operator Dashboard")
    add_bullet("My Shift Completion Counter: Shows number of checklists completed today vs scheduled targets.", "• ")
    add_bullet("Defect Rate / Open Issues: Highlights any open non-conformances logged during the current shift.", "• ")
    add_bullet("Scheduled Inspections: Live cards indicating upcoming inspections due within the shift.", "• ")

    add_h2("10.2 Supervisor Dashboard")
    add_bullet("Line-by-Line Inspection Progress: Real-time progress bars for all lines in the shop floor.", "• ")
    add_bullet("Missed Inspections Tracker: Alerts supervisors if a line missed an hourly inspection interval.", "• ")
    add_bullet("Top Defect Categories: Bar charts showing which components or stations have the highest defect counts.", "• ")
    add_bullet("Team Performance Summary: Overview of operator submissions and response velocity.", "• ")

    add_h2("10.3 Management Dashboard")
    add_bullet("Plant-Wide First-Pass Quality Rate: Aggregated quality trend across all shops and departments.", "• ")
    add_bullet("Cross-Shop Comparison: Side-by-side performance benchmarks for SMT, Assembly, Testing, and Packaging.", "• ")
    add_bullet("CAPA Resolution Health: Number of Open, In-Progress, and Overdue CAPA tickets across the facility.", "• ")
    add_bullet("ISO / Customer Audit Readiness: Instant exportable compliance reports for management reviews.", "• ")

    # =========================================================================
    # SECTION 11: INSPECTION SCHEDULER
    # =========================================================================
    add_h1("11. Inspection Scheduler & Shift Routine Triggers")
    add_body(
        "The Inspection Scheduler automates quality routines so operators never miss mandatory inspections. "
        "Quality administrators and supervisors configure how often each checklist must be executed."
    )

    add_h2("11.1 Available Schedule Types")
    add_bullet("Interval-Based (e.g., Every 60 Minutes): Automatically prompts operators to perform an inspection at fixed minute intervals throughout the shift.", "1. Interval Schedule: ")
    add_bullet("Shift-Based (e.g., 3 Times Per Shift): Automatically splits the shift into equal windows and tracks completion per window (e.g., Start of Shift, Mid-Shift, End of Shift).", "2. Shift Limit Schedule: ")
    add_bullet("Daily Fixed Time (e.g., Every day at 08:30 AM): Automatically triggers daily morning line verification or calibration checks at a specific hour.", "3. Daily Schedule: ")

    add_h2("11.2 Step-by-Step: Setting Up an Inspection Schedule")
    add_step("1", "Open Scheduler Admin", "Navigate to Scheduler from the admin menu.")
    add_step("2", "Click [Create Schedule]", "Click the button to add a new scheduling rule.")
    add_step("3", "Select Checklist Template", "Choose the quality form you want to schedule.")
    add_step("4", "Select Schedule Type & Interval", "Pick Interval, Times-Per-Shift, or Daily Fixed Time and input the value (e.g., 60 minutes).")
    add_step("5", "Set Status to Active", "Toggle the status to 'Active' and click [Save Schedule]. The system will now automatically track compliance and alert supervisors if an interval is missed.")

    # =========================================================================
    # SECTION 12: FORM BUILDER GUIDE (FOR QUALITY ENGINEERS)
    # =========================================================================
    add_h1("12. Form Builder Guide (Checklist & Dynamic Form Setup)")
    add_body(
        "Quality Engineers can create and update inspection forms digitally without writing code, "
        "complete with version control and ISO compliance tracking."
    )

    add_h2("12.1 Step-by-Step: Creating a New Checklist Template")
    add_step("1", "Access Forms Builder", "Navigate to TranssFlow Forms Builder or Dynamic Forms Admin.")
    add_step("2", "Click [Create New Template]", "Enter Form Name (e.g., 'Final Quality Inspection'), Form Code (e.g., 'FQC-CHK-001'), and Description.")
    add_step("3", "Add Inspection Sections", "Click 'Add Section' to organize questions logically (e.g., '1. Visual Appearance', '2. Dimensional Checks', '3. Functional Testing').")
    add_step("4", "Add Inspection Items", "Under each section, click 'Add New Item' and define: Item Label (question text), Input Type (Yes/No, Number, Text, Photo, Dropdown), Severity Weight (0 for minor, 5 for critical), and whether the item is Required.")
    add_step("5", "Configure Conditional Rules", "Select 'Photo Required on Fail' and 'Remark Required on Fail' to automatically enforce evidence collection when defects occur.")
    add_step("6", "Preview & Finalize Version", "Click 'Preview' to test how the form looks on mobile/tablet. When satisfied, click 'Finalize Template Version' to publish it to the shop floor.")

    # =========================================================================
    # SECTION 13: PLANT ADMINISTRATION & USER MANAGEMENT
    # =========================================================================
    add_h1("13. Plant Administration & User Management")
    add_body("Plant Administrators configure organizational structures and manage user accounts in the Admin Control Center (`/admin_panel/`).")

    add_h2("13.1 Managing Organization Structure")
    add_bullet("Plants: Configure manufacturing facilities (e.g., Plant 1, Plant 2).", "• ")
    add_bullet("Shops & Departments: Define shop floor areas (e.g., SMT, Assembly, Quality, Maintenance).", "• ")
    add_bullet("Production Lines: Add and name production lines with unique line codes.", "• ")
    add_bullet("Stations: Add specific workstations along each line (e.g., Station 01 - Soldering, Station 05 - Packaging).", "• ")
    add_bullet("Products / Models: Add active manufactured models, part numbers, and color variants.", "• ")

    add_h2("13.2 User Management & Scope Assignment")
    add_step("1", "Navigate to Accounts > User Management", "Open the user list in the Admin Panel.")
    add_step("2", "Create New User or Bulk Upload", "Click 'Create User' to add an individual employee, or click 'Bulk Upload via Excel' to import multiple operators at once.")
    add_step("3", "Assign Plant & Line Scope", "Assign the user to their specific Plant and Production Line. This ensures operators only see forms relevant to their work area.")
    add_step("4", "Assign Role", "Select the appropriate role (Operator, Supervisor, PQE, EHS Auditor, Manager).")

    # =========================================================================
    # SECTION 14: OFFLINE MODE & AUTO-SYNC GUIDELINES
    # =========================================================================
    add_h1("14. Offline Mode & Auto-Sync Guidelines")
    add_body(
        "QIsmartuX includes built-in Progressive Web App (PWA) offline capabilities. "
        "If plant Wi-Fi temporarily drops while an operator is on the shop floor, work is not lost."
    )

    add_callout_box(
        doc,
        "How Offline Mode Works for Shop Floor Users",
        "1. When network connection is lost, an 'Offline Mode' badge appears in the top corner.\n"
        "2. Operators can continue filling in checklist answers and taking photos as normal.\n"
        "3. Responses and photos are safely stored in the tablet's local secure memory.\n"
        "4. As soon as Wi-Fi reconnects, the system automatically syncs all offline submissions to the central server without data loss.",
        "tip"
    )

    # =========================================================================
    # SECTION 15: BEST PRACTICES, DO'S & DON'TS, FAQS & TROUBLESHOOTING
    # =========================================================================
    add_h1("15. Best Practices, Do's & Don'ts, FAQs and Troubleshooting")

    add_h2("15.1 Daily Operator Checklist (Shift Routine)")
    add_step("1", "Shift Start (0 - 15 Mins)", "Log into QMS > Check/Activate Work Session (Line, Model, Shift) > Review Scheduled Inspections.")
    add_step("2", "During Shift Operations", "Perform scheduled inspections on time > Log exact measurements > If a defect occurs, mark 'Fail', take a clear photo, and type the exact reason.")
    add_step("3", "Shift Handover / End of Shift", "Verify all shift checklists are submitted > Confirm no pending drafts > Hand over active status to incoming shift operator > Log out.")

    add_h2("15.2 Quality Do's and Don'ts")
    
    dodont_table = doc.add_table(rows=6, cols=2)
    dodont_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    dodont_table.autofit = False
    set_table_borders(dodont_table, "CBD5E1", "4")
    
    dd_widths = [Inches(3.25), Inches(3.25)]
    for row in dodont_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = dd_widths[i]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # Header Row
    headers_dd = ["DO'S (Mandatory Practices)", "DON'TS (Strictly Prohibited)"]
    for i, h in enumerate(headers_dd):
        cell = dodont_table.cell(0, i)
        set_cell_background(cell, "16A34A" if i == 0 else "DC2626")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    dd_data = [
        ("Always verify your active line and model before starting an inspection.", "Never perform inspections under another operator's logged-in session."),
        ("Capture sharp, well-lit photographs focused directly on the defect area.", "Never upload blurry, dark, or placeholder photos for failed items."),
        ("Enter accurate numerical measurements as read directly from calibrated gauges.", "Never estimate or round off critical parameter measurements."),
        ("Submit inspections immediately after performing the physical check.", "Never batch or delay submissions to the end of the shift."),
        ("Report repeated defects immediately to your line leader and raise a CAPA.", "Never bypass or ignore red validation warnings on the inspection screen.")
    ]

    for row_idx, data in enumerate(dd_data, start=1):
        bg = "F0FDF4" if row_idx % 2 == 1 else "FEF2F2"
        for col_idx in range(2):
            cell = dodont_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(data[col_idx])
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            run.font.color.rgb = SLATE_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("15.3 Frequently Asked Questions (FAQs)")
    
    faqs = [
        ("Q1: Why is the 'Submit Inspection' button not working or showing red?",
         "A1: This happens when one or more required questions have not been answered, or an item marked 'Fail' is missing its mandatory defect photo or remark. Scroll up and look for the red 'Required field missing' alerts, complete them, and try submitting again."),
        
        ("Q2: How do I change the model if the line switches products during the shift?",
         "A2: Click the 'Change Session' button at the top right of your screen. Select the new Product / Model from the dropdown list, select your shift, and click 'Activate Session'. All subsequent inspections will be tagged to the new model."),
         
        ("Q3: What should I do if my tablet camera does not open when uploading photos?",
         "A3: Ensure browser camera permissions are enabled. Tap the lock icon in your browser's address bar, choose 'Site Settings' or 'Permissions', and set 'Camera' to 'Allow'. Then refresh the inspection page."),
         
        ("Q4: Can I edit an inspection after it has been submitted?",
         "A4: No. For quality integrity and ISO traceability, submitted inspection records are locked. If an error was made, inform your supervisor or PQE engineer. They can reject the submission with comments, allowing re-inspection."),
         
        ("Q5: Where can I see which inspections I still need to complete today?",
         "A5: Visit your 'Operator Dashboard' or the 'Forms List' page. Scheduled inspection cards will indicate how many checks are completed vs remaining for your current shift.")
    ]

    for q, a in faqs:
        add_body(a, bold_prefix=f"{q}\n")

    add_h2("15.4 Troubleshooting Quick-Fix Table")
    
    tb_table = doc.add_table(rows=5, cols=3)
    tb_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb_table.autofit = False
    set_table_borders(tb_table, "CBD5E1", "4")
    
    t_widths = [Inches(1.8), Inches(2.2), Inches(2.5)]
    for row in tb_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = t_widths[i]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # Header Row
    headers_tb = ["Symptom / Problem", "Possible Root Cause", "Step-by-Step Resolution"]
    for i, h in enumerate(headers_tb):
        cell = tb_table.cell(0, i)
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    tb_data = [
        ("No forms appear on the checklist page.", "Work session not active or line scope not assigned.", "Click 'Change Session' and activate your Line and Model. If still empty, ask admin to assign the template to your line."),
        ("Red validation error on submit.", "Unanswered mandatory question or missing defect photo.", "Review form sections. Check all items with a red asterisk (*) or failed items requiring a photo/remark."),
        ("'No Role Access Configured' message.", "User account created but role permissions pending.", "Contact Plant System Administrator to assign your role (Operator, Supervisor, PQE, or EHS)."),
        ("Submission stuck on 'Syncing'.", "Weak or disconnected shop floor Wi-Fi.", "Check device Wi-Fi connection. The system will auto-sync once connection stabilizes.")
    ]

    for row_idx, data in enumerate(tb_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(3):
            cell = tb_table.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(data[col_idx])
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = SLATE_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Final Sign-off block
    sign_table = doc.add_table(rows=2, cols=3)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    set_table_borders(sign_table, "CBD5E1", "4")
    for row in sign_table.rows:
        for cell in row.cells:
            cell.width = Inches(2.16)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    sign_headers = ["Prepared By:", "Reviewed By:", "Approved By:"]
    for i, h in enumerate(sign_headers):
        cell = sign_table.cell(0, i)
        set_cell_background(cell, "F1F5F9")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(71, 85, 105)

    sign_data = [
        "Quality Operations Team\nDate: 24-Aug-2026",
        "Lead Quality Engineer (PQE)\nDate: 24-Aug-2026",
        "Plant Operations & Quality Head\nDate: 24-Aug-2026"
    ]
    for i, d in enumerate(sign_data):
        cell = sign_table.cell(1, i)
        p = cell.paragraphs[0]
        run = p.add_run(d)
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = SLATE_BODY

    output_path = os.path.join(os.getcwd(), "QMS_User_Standard_Operating_Procedure_SOP.docx")
    doc.save(output_path)
    print(f"Successfully generated SOP DOCX at: {output_path}")

if __name__ == "__main__":
    build_sop()
